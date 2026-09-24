"""
Automated tests for Cloud Web OAuth 2.0 flow and serverless session handling.
Tests /api/auth/* endpoints, token encryption, and serverless Excel streaming.
Does NOT require real Google credentials (uses mocks).
"""

import io
import json
import os
from unittest.mock import MagicMock, patch
from fastapi.testclient import TestClient
import openpyxl
import pytest

from app.api import app
from app.models import AttendanceRecord, AttendanceStatus, ProcessingLogEntry
from app.oauth import (
    create_authorization_url,
    decrypt_session_data,
    encrypt_session_data,
    exchange_code_for_session,
    get_credentials_from_session_token,
)
from app.excel_writer import export_attendance_bytes

client = TestClient(app)


@pytest.fixture(autouse=True)
def mock_oauth_env(monkeypatch):
    """Sets standard mock OAuth environment variables for cloud tests."""
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "test-client-id-12345.apps.googleusercontent.com")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "test-client-secret-abcde")
    monkeypatch.setenv("GOOGLE_REDIRECT_URI", "https://test-belvo.vercel.app/api/auth/callback")
    monkeypatch.setenv("SESSION_SECRET", "test-session-secret-key-for-fernet-32b")


class TestCloudOAuthFlow:
    """Tests for Web OAuth 2.0 flow on Vercel."""

    def test_auth_google_json_response(self):
        """Verify /api/auth/google returns valid JSON with authorization URL."""
        response = client.get("/api/auth/google?json_response=true")
        assert response.status_code == 200
        data = response.json()
        assert "auth_url" in data
        assert "state" in data

        auth_url = data["auth_url"]
        assert "accounts.google.com/o/oauth2/v2/auth" in auth_url
        assert "test-client-id-12345" in auth_url
        assert "gmail.readonly" in auth_url
        assert "access_type=offline" in auth_url
        assert "prompt=consent" in auth_url

    def test_auth_google_redirect(self):
        """Verify /api/auth/google redirects browser by default."""
        response = client.get("/api/auth/google", follow_redirects=False)
        assert response.status_code == 307
        assert "accounts.google.com" in response.headers["location"]

    def test_auth_callback_missing_code(self):
        """Verify /api/auth/callback rejects missing code."""
        response = client.get("/api/auth/callback")
        assert response.status_code == 400
        assert "Missing authorization code" in response.json()["detail"]

    def test_auth_callback_error_param(self):
        """Verify /api/auth/callback handles Google error response."""
        response = client.get("/api/auth/callback?error=access_denied")
        assert response.status_code == 400
        assert "access_denied" in response.json()["detail"]

    @patch("urllib.request.urlopen")
    def test_auth_callback_successful_exchange(self, mock_urlopen):
        """Verify successful code exchange generates session token and HTML response."""
        # Mock Google Token endpoint response
        mock_response = MagicMock()
        mock_response.read.return_value = json.dumps({
            "access_token": "ya29.mock-access-token-123",
            "refresh_token": "1//mock-refresh-token-456",
            "expires_in": 3599,
            "token_type": "Bearer",
            "scope": "https://www.googleapis.com/auth/gmail.readonly",
        }).encode("utf-8")
        mock_urlopen.return_value.__enter__.return_value = mock_response

        response = client.get("/api/auth/callback?code=mock-google-auth-code&state=mock-state")
        assert response.status_code == 200
        assert "text/html" in response.headers["content-type"]
        assert "Gmail Connected Successfully!" in response.text
        assert "set-cookie" in response.headers

    def test_session_token_encryption_and_decryption(self):
        """Verify symmetric encryption protects OAuth credentials."""
        test_payload = {
            "token": "secret-access-token",
            "refresh_token": "secret-refresh-token-xyz",
            "client_id": "test-client",
            "client_secret": "test-secret",
        }

        # Encrypt
        encrypted = encrypt_session_data(test_payload)
        assert isinstance(encrypted, str)
        assert "secret-refresh-token-xyz" not in encrypted  # Never stored in plaintext

        # Decrypt
        decrypted = decrypt_session_data(encrypted)
        assert decrypted == test_payload

    def test_auth_status_with_valid_and_invalid_token(self):
        """Verify /api/auth/status checks session token validity."""
        # Invalid token
        res_invalid = client.get("/api/auth/status", headers={"Authorization": "Bearer invalid.token"})
        assert res_invalid.status_code == 200
        assert res_invalid.json()["authenticated"] is False

        # Valid encrypted session token
        valid_creds = {
            "token": "valid-token-123",
            "refresh_token": "valid-refresh-456",
            "token_uri": "https://oauth2.googleapis.com/token",
            "client_id": "test-client-id",
            "client_secret": "test-secret",
            "scopes": ["https://www.googleapis.com/auth/gmail.readonly"],
        }
        valid_session = encrypt_session_data(valid_creds)

        res_valid = client.get(
            "/api/auth/status",
            headers={"Authorization": f"Bearer {valid_session}"},
        )
        assert res_valid.status_code == 200
        assert res_valid.json()["authenticated"] is True
        assert res_valid.json()["method"] == "web_oauth"

    def test_auth_logout(self):
        """Verify /api/auth/logout endpoint."""
        response = client.post("/api/auth/logout")
        assert response.status_code == 200
        assert response.json()["authenticated"] is False


class TestServerlessExcelBytes:
    """Tests for in-memory Excel generation on serverless environments."""

    def test_export_attendance_bytes(self):
        records = [
            AttendanceRecord(
                date="2026-09-17",
                person="Pavan Atchyuta",
                email="pavan@example.com",
                status=AttendanceStatus.PRESENT,
            ),
            AttendanceRecord(
                date="2026-09-17",
                person="Ananya",
                email="ananya@example.com",
                status=AttendanceStatus.LEAVE,
            ),
        ]
        logs = [
            ProcessingLogEntry(
                timestamp="2026-09-17T18:00:00Z",
                category=AttendanceStatus.PRESENT,
                sender="pavan@example.com",
                subject="Daily Work Report - 2026-09-17",
                target_date="2026-09-17",
                action="Accepted",
                details="Valid report",
            )
        ]

        excel_bytes = export_attendance_bytes(records, logs)
        assert isinstance(excel_bytes, bytes)
        assert len(excel_bytes) > 0

        # Verify openpyxl can read the bytes
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert "Attendance" in wb.sheetnames
        assert "Processing Log" in wb.sheetnames

        ws = wb["Attendance"]
        assert ws.cell(row=2, column=2).value == "Pavan Atchyuta"
        assert ws.cell(row=2, column=4).value == "P"
        assert ws.cell(row=3, column=4).value == "L"
