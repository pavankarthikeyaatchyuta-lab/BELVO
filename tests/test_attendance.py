"""
Comprehensive unit and integration test suite for Belvo Attendance Tracker.
Covers all required edge cases and acceptance criteria.
"""

from pathlib import Path
import tempfile
import openpyxl
import pytest

from app.attendance import AttendanceEngine, load_employees_from_csv, load_leave_from_csv
from app.excel_writer import export_attendance_workbook
from app.models import (
    AttendanceRecord,
    AttendanceStatus,
    EmailMessage,
    Employee,
    LeaveEntry,
    LogCategory,
)
from app.parser import extract_clean_email, parse_work_report


@pytest.fixture
def sample_employees():
    return [
        Employee(name="Pavan Atchyuta", email="pavan@example.com"),
        Employee(name="Thrija", email="thrija@example.com"),
        Employee(name="Sai", email="sai@example.com"),
        Employee(name="Ravi", email="ravi@example.com"),
        Employee(name="Ananya", email="ananya@example.com"),
    ]


@pytest.fixture
def sample_leave():
    return [
        LeaveEntry(email="ananya@example.com", date="2026-09-17"),
    ]


class TestEmailParser:
    """Tests for email normalization and subject parsing."""

    def test_extract_clean_email(self):
        assert extract_clean_email("Pavan Atchyuta <pavan@example.com>") == "pavan@example.com"
        assert extract_clean_email("<THRIJA@EXAMPLE.COM> ") == "thrija@example.com"
        assert extract_clean_email("  sai@example.com  ") == "sai@example.com"
        assert extract_clean_email("") == ""
        assert extract_clean_email(None) == ""

    def test_valid_subject_formats(self):
        valid_subjects = [
            "Daily Work Report - 2026-09-17",
            "daily work report - 2026-09-17",
            "Daily Work Report: 2026-09-17",
            "  Daily Work Report - 2026-09-17  ",
            "daily work report: 2026-09-17",
        ]
        for subj in valid_subjects:
            msg = EmailMessage(id="1", sender="user@example.com", subject=subj)
            parsed = parse_work_report(msg, "2026-09-17")
            assert parsed.is_valid is True, f"Failed for subject: {subj}"
            assert parsed.extracted_date == "2026-09-17"

    def test_malformed_subjects(self):
        invalid_subjects = [
            "Hello",
            "Work",
            "Daily update",
            "Report",
            "Daily Work Report",
            "Daily Work Report - yesterday",
            "Daily Work Report - 2026-99-99",  # Invalid calendar date
            "Daily Work Report - 17-09-2026",  # Wrong date format
        ]
        for subj in invalid_subjects:
            msg = EmailMessage(id="1", sender="user@example.com", subject=subj)
            parsed = parse_work_report(msg, "2026-09-17")
            assert parsed.is_valid is False
            assert parsed.category == LogCategory.MALFORMED_SUBJECT

    def test_missing_subject(self):
        for empty_subj in [None, "", "   "]:
            msg = EmailMessage(id="1", sender="user@example.com", subject=empty_subj)
            parsed = parse_work_report(msg, "2026-09-17")
            assert parsed.is_valid is False
            assert parsed.category == LogCategory.MISSING_SUBJECT


class TestAttendanceEngine:
    """Tests for attendance calculation logic."""

    def test_normal_present(self, sample_employees):
        messages = [
            EmailMessage(
                id="1",
                sender="Pavan Atchyuta <pavan@example.com>",
                subject="Daily Work Report - 2026-09-17",
                received_at="2026-09-17T18:00:00+00:00",
            )
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        pavan_rec = next(r for r in records if r.email == "pavan@example.com")
        assert pavan_rec.status == AttendanceStatus.PRESENT
        assert stats["present_count"] == 1
        assert stats["absent_count"] == 4

    def test_absent(self, sample_employees):
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance([], "2026-09-17")

        assert len(records) == 5
        assert all(r.status == AttendanceStatus.ABSENT for r in records)
        assert stats["absent_count"] == 5

    def test_leave_status(self, sample_employees, sample_leave):
        engine = AttendanceEngine(employees=sample_employees, leave_entries=sample_leave)
        records, logs, stats = engine.process_attendance([], "2026-09-17")

        ananya_rec = next(r for r in records if r.email == "ananya@example.com")
        assert ananya_rec.status == AttendanceStatus.LEAVE
        assert stats["leave_count"] == 1

    def test_leave_precedence_over_report(self, sample_employees, sample_leave):
        """If person is on leave AND submitted a report, leave takes precedence."""
        messages = [
            EmailMessage(
                id="1",
                sender="Ananya <ananya@example.com>",
                subject="Daily Work Report - 2026-09-17",
                received_at="2026-09-17T16:00:00+00:00",
            )
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=sample_leave)
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        ananya_rec = next(r for r in records if r.email == "ananya@example.com")
        assert ananya_rec.status == AttendanceStatus.LEAVE
        assert "leave precedence" in ananya_rec.notes

    def test_duplicate_reports_deduplicated(self, sample_employees):
        """Multiple reports from same person for same day must produce only ONE attendance row."""
        messages = [
            EmailMessage(
                id="1",
                sender="pavan@example.com",
                subject="Daily Work Report - 2026-09-17",
                received_at="2026-09-17T18:00:00+00:00",
            ),
            EmailMessage(
                id="2",
                sender="pavan@example.com",
                subject="Daily Work Report - 2026-09-17",
                received_at="2026-09-17T18:10:00+00:00",
            ),
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        pavan_records = [r for r in records if r.email == "pavan@example.com"]
        assert len(pavan_records) == 1
        assert pavan_records[0].status == AttendanceStatus.PRESENT
        assert stats["duplicate_reports"] == 1

    def test_late_report_handling(self, sample_employees):
        """Report received on 2026-09-18 01:15 AM with subject 2026-09-17 counts for 2026-09-17."""
        messages = [
            EmailMessage(
                id="1",
                sender="Sai <sai@example.com>",
                subject="Daily Work Report - 2026-09-17",
                received_at="2026-09-18T01:15:00+00:00",
            )
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        sai_rec = next(r for r in records if r.email == "sai@example.com")
        assert sai_rec.status == AttendanceStatus.PRESENT
        assert stats["late_reports"] == 1

    def test_malformed_and_missing_subjects(self, sample_employees):
        messages = [
            EmailMessage(id="1", sender="ravi@example.com", subject="Hello team"),
            EmailMessage(id="2", sender="ravi@example.com", subject=""),
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        ravi_rec = next(r for r in records if r.email == "ravi@example.com")
        assert ravi_rec.status == AttendanceStatus.ABSENT
        assert stats["malformed_subjects"] == 1
        assert stats["missing_subjects"] == 1

    def test_unknown_sender(self, sample_employees):
        """Unknown senders must not be added to employee roster or mark anyone present."""
        messages = [
            EmailMessage(
                id="1",
                sender="stranger@outsider.com",
                subject="Daily Work Report - 2026-09-17",
            )
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        assert len(records) == len(sample_employees)
        assert not any(r.email == "stranger@outsider.com" for r in records)
        assert stats["unknown_senders"] == 1
        assert any(l.category == LogCategory.UNKNOWN_SENDER for l in logs)

    def test_case_insensitive_matching(self, sample_employees):
        messages = [
            EmailMessage(
                id="1",
                sender="PAVAN@EXAMPLE.COM",
                subject="daily work report - 2026-09-17",
            )
        ]
        engine = AttendanceEngine(employees=sample_employees, leave_entries=[])
        records, logs, stats = engine.process_attendance(messages, "2026-09-17")

        pavan_rec = next(r for r in records if r.email == "pavan@example.com")
        assert pavan_rec.status == AttendanceStatus.PRESENT


class TestExcelWriter:
    """Tests for Excel workbook generation."""

    def test_excel_export_structure(self, sample_employees):
        records = [
            AttendanceRecord(
                date="2026-09-17",
                person="Pavan Atchyuta",
                email="pavan@example.com",
                status=AttendanceStatus.PRESENT,
            ),
            AttendanceRecord(
                date="2026-09-17",
                person="Ravi",
                email="ravi@example.com",
                status=AttendanceStatus.ABSENT,
            ),
            AttendanceRecord(
                date="2026-09-17",
                person="Ananya",
                email="ananya@example.com",
                status=AttendanceStatus.LEAVE,
            ),
        ]
        logs = []

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_attendance.xlsx"
            export_attendance_workbook(records, logs, out_file)

            assert out_file.exists()

            wb = openpyxl.load_workbook(out_file)
            assert "Attendance" in wb.sheetnames
            assert "Processing Log" in wb.sheetnames

            ws = wb["Attendance"]
            assert ws.cell(row=1, column=1).value == "Date"
            assert ws.cell(row=1, column=2).value == "Person"
            assert ws.cell(row=1, column=3).value == "Email"
            assert ws.cell(row=1, column=4).value == "Status"

            # Check rows
            assert ws.cell(row=2, column=4).value == "P"
            assert ws.cell(row=3, column=4).value == "A"
            assert ws.cell(row=4, column=4).value == "L"
