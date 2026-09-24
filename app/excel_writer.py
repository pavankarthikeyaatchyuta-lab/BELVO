"""
Excel report generation module using openpyxl.
Generates styled 'Attendance' worksheet and auditable 'Processing Log' worksheet.
"""

from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import AttendanceRecord, AttendanceStatus, ProcessingLogEntry


# Theme styling palette
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

STATUS_STYLES = {
    AttendanceStatus.PRESENT: {
        "fill": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="155724"),
    },
    AttendanceStatus.ABSENT: {
        "fill": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="721C24"),
    },
    AttendanceStatus.LEAVE: {
        "fill": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="856404"),
    },
}

REGULAR_FONT = Font(name="Calibri", size=11, color="000000")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)


import io

def build_attendance_workbook(
    attendance_records: List[AttendanceRecord],
    processing_logs: List[ProcessingLogEntry],
) -> openpyxl.Workbook:
    """
    Constructs and styles an openpyxl Workbook with 'Attendance' and 'Processing Log' sheets.
    """
    wb = openpyxl.Workbook()

    # --- Sheet 1: Attendance ---
    ws_att = wb.active
    ws_att.title = "Attendance"
    ws_att.views.sheetView[0].showGridLines = True

    att_headers = ["Date", "Person", "Email", "Status"]
    ws_att.append(att_headers)

    # Style Header
    for col_idx, header in enumerate(att_headers, start=1):
        cell = ws_att.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    ws_att.freeze_panes = "A2"
    ws_att.row_dimensions[1].height = 26

    # Populate Attendance Rows
    for row_idx, rec in enumerate(attendance_records, start=2):
        ws_att.row_dimensions[row_idx].height = 22

        c_date = ws_att.cell(row=row_idx, column=1, value=rec.date)
        c_person = ws_att.cell(row=row_idx, column=2, value=rec.person)
        c_email = ws_att.cell(row=row_idx, column=3, value=rec.email)
        c_status = ws_att.cell(row=row_idx, column=4, value=rec.status.value)

        c_date.alignment = CENTER_ALIGN
        c_date.font = REGULAR_FONT
        c_date.border = THIN_BORDER

        c_person.alignment = LEFT_ALIGN
        c_person.font = REGULAR_FONT
        c_person.border = THIN_BORDER

        c_email.alignment = LEFT_ALIGN
        c_email.font = REGULAR_FONT
        c_email.border = THIN_BORDER

        # Status cell with specific fill and font
        c_status.alignment = CENTER_ALIGN
        c_status.border = THIN_BORDER
        style_info = STATUS_STYLES.get(rec.status)
        if style_info:
            c_status.fill = style_info["fill"]
            c_status.font = style_info["font"]
        else:
            c_status.font = REGULAR_FONT

    # Adjust Attendance Column Widths
    for col in ws_att.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_att.column_dimensions[col_letter].width = max(max_len + 5, 14)

    # --- Sheet 2: Processing Log ---
    ws_log = wb.create_sheet(title="Processing Log")
    ws_log.views.sheetView[0].showGridLines = True

    log_headers = ["Timestamp", "Category", "Sender", "Subject", "Target Date", "Action", "Details"]
    ws_log.append(log_headers)

    for col_idx, header in enumerate(log_headers, start=1):
        cell = ws_log.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    ws_log.freeze_panes = "A2"
    ws_log.row_dimensions[1].height = 26

    for row_idx, log_entry in enumerate(processing_logs, start=2):
        ws_log.row_dimensions[row_idx].height = 20
        row_values = [
            log_entry.timestamp,
            log_entry.category.value if hasattr(log_entry.category, "value") else str(log_entry.category),
            log_entry.sender,
            log_entry.subject,
            log_entry.target_date,
            log_entry.action,
            log_entry.details,
        ]
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws_log.cell(row=row_idx, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_ALIGN if col_idx in (1, 2, 5, 6) else LEFT_ALIGN

    # Adjust Log Column Widths
    for col in ws_log.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_log.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

    return wb


def export_attendance_workbook(
    attendance_records: List[AttendanceRecord],
    processing_logs: List[ProcessingLogEntry],
    output_path: Path,
) -> Path:
    """
    Creates and saves an Excel workbook to a disk path.
    Used by CLI, local development, and tests.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_attendance_workbook(attendance_records, processing_logs)
    wb.save(str(output_path))
    return output_path


def export_attendance_bytes(
    attendance_records: List[AttendanceRecord],
    processing_logs: List[ProcessingLogEntry],
) -> bytes:
    """
    Creates an Excel workbook in memory and returns raw bytes.
    Ideal for serverless environments (Vercel) without persistent disk.
    """
    wb = build_attendance_workbook(attendance_records, processing_logs)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

